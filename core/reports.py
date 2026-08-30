"""PDF and Excel attendance reports."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core import db
from core.time_utils import period_range

RED = PatternFill("solid", fgColor="FFC7C7")
GREEN = PatternFill("solid", fgColor="C6F4D6")
PURPLE = PatternFill("solid", fgColor="E3D0F2")
PDF_RED = colors.HexColor("#FFC7C7")
PDF_GREEN = colors.HexColor("#C6F4D6")
PDF_PURPLE = colors.HexColor("#E3D0F2")
HEADERS = ["Name", "ID", "Shift", "Clock in", "Clock out", "Total hours", "Late"]


def _fill_for(status: str):
    if status == "under":
        return RED, PDF_RED
    if status == "exact":
        return GREEN, PDF_GREEN
    if status == "over":
        return PURPLE, PDF_PURPLE
    return None, None


def resolve_period(period: str, work_date: str) -> tuple[str, str, str]:
    day = date.fromisoformat(work_date) if work_date.strip() else date.today()
    start, end = period_range(period, day)
    return start.isoformat(), end.isoformat(), (period or "daily").lower()


def report_rows(period: str, work_date: str) -> tuple[list[dict], str, str, str]:
    db.init_db()
    start, end, kind = resolve_period(period, work_date)
    return db.list_sessions(date_from=start, date_to=end), start, end, kind


def _row_values(r: dict) -> list:
    return [
        r["employee_name"],
        r["employee_id"],
        r.get("shift_name") or "—",
        _fmt_time(r.get("clock_in")),
        _fmt_time(r.get("clock_out")) if r.get("clock_out") else "—",
        r["hours_formatted"],
        r.get("late_label") or "On time",
    ]


def build_excel(period: str, work_date: str) -> bytes:
    rows, start, end, kind = report_rows(period, work_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append([f"ACCESS PASS {kind} report", f"{start} to {end}"])
    ws.append(HEADERS)
    for cell in ws[2]:
        cell.font = Font(bold=True)

    for r in rows:
        ws.append(_row_values(r))
        excel_fill, _ = _fill_for(r.get("hours_status") or "")
        if excel_fill and r.get("clock_out"):
            for cell in ws[ws.max_row]:
                cell.fill = excel_fill

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
        for cell in col:
            cell.alignment = Alignment(vertical="center")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(period: str, work_date: str) -> bytes:
    rows, start, end, kind = report_rows(period, work_date)
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=f"ACCESS PASS {kind}")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("ACCESS PASS", styles["Title"]),
        Paragraph(
            f"{kind.title()} attendance · {start} to {end} · late uses actual hours",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    data = [HEADERS]
    status_by_row = [None]
    for r in rows:
        data.append(_row_values(r))
        status_by_row.append(r.get("hours_status") if r.get("clock_out") else None)

    table = Table(data, repeatRows=1)
    style_cmds = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.black),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]
    for i, status in enumerate(status_by_row):
        _, pdf_fill = _fill_for(status or "")
        if pdf_fill:
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), pdf_fill))
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def _fmt_time(iso: str | None) -> str:
    if not iso:
        return "—"
    try:
        return datetime_from(iso).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return iso


def datetime_from(iso: str):
    from datetime import datetime as dt

    return dt.fromisoformat(iso)
